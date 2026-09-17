# -*- coding: utf-8 -*-
"""将新知识图谱图替换进 v23（图13/14/A1-A3），并同步显示尺寸防变形。
替换映射（media ← 新图）：
  image13.png ← _fig/fig13_kg_overall.png     图13 主图（整体）   2109×2109
  image14.png ← _fig/fig_kg_grid_acl.png      图14 六视图网格     3800×2500
  image16.png ← _fig/kg_acl_view1_overall.png 图A1 整体视图      1878×1878
  image17.png ← _fig/kg_acl_view4_zoom.png    图A2 核心连接簇放大 1878×1878
  image18.png ← _fig/kg_acl_view6_closeup.png 图A3 RADAR近景     1878×1878
显示尺寸按新图比例（cm→EMU, ×360000）：
  图13 14.5×14.5 / 图14 14.5×9.54 / 图A1 13.5×13.5 / 图A2 11.7×11.7 / 图A3 11.7×11.7
完整版输出 v24（v23 被占用无法原地覆盖，避免产生不完整中间版）。
"""
import sys, os, zipfile, time
sys.stdout.reconfigure(encoding='utf-8')
from docx import Document
from docx.oxml.ns import qn

SRC = '论文初稿_new_v23_中文版_final13_图.docx'
DST = '论文初稿_new_v24_中文版_final13_图.docx'
EMU = 360000
NS_A = '{http://schemas.openxmlformats.org/drawingml/2006/main}'
NS_R = '{http://schemas.openxmlformats.org/officeDocument/2006/relationships}'
NS_WP = '{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}'

# media 名 → (新图路径, 新显示尺寸 cm 宽, 高)
PLAN = {
    'image13.png': ('_fig/fig13_kg_overall.png',      14.5, 14.5),
    'image14.png': ('_fig/fig_kg_grid_acl.png',       14.5, 9.54),
    'image16.png': ('_fig/kg_acl_view1_overall.png',  13.5, 13.5),
    'image17.png': ('_fig/kg_acl_view4_zoom.png',     11.7, 11.7),
    'image18.png': ('_fig/kg_acl_view6_closeup.png',  11.7, 11.7),
}

# ---------- 1. python-docx 更新显示尺寸 → 保存到 v24 ----------
doc = Document(SRC)
updated = []
for p in doc.paragraphs:
    for d in p._p.iter(qn('w:drawing')):
        blip = d.find('.//' + NS_A + 'blip')
        if blip is None:
            continue
        rId = blip.get(NS_R + 'embed')
        part = doc.part.related_parts.get(rId)
        media = part.partname.split('/')[-1] if part is not None else None
        if media not in PLAN:
            continue
        cx = int(PLAN[media][1] * EMU)
        cy = int(PLAN[media][2] * EMU)
        # wp:inline/wp:extent
        for ext in d.iter(NS_WP + 'extent'):
            ext.set('cx', str(cx)); ext.set('cy', str(cy))
        # a:xfrm/a:ext
        for ext in d.iter(NS_A + 'ext'):
            if ext.get('cx') is not None:
                ext.set('cx', str(cx)); ext.set('cy', str(cy))
        updated.append((media, PLAN[media][1], PLAN[media][2]))
doc.save(DST)
print('✓ 显示尺寸已更新:')
for m, w, h in updated:
    print(f'   {m} → {w}×{h}cm')

# ---------- 2. zipfile 替换 media 字节（v24，带重试防杀毒瞬时锁） ----------
TMP = '论文初稿_new_v24_tmp.docx'
if os.path.exists(TMP):
    os.remove(TMP)
REPLACE = {('word/media/' + k): v[0] for k, v in PLAN.items()}
with zipfile.ZipFile(DST, 'r') as zin, zipfile.ZipFile(TMP, 'w', zipfile.ZIP_DEFLATED) as zout:
    for item in zin.infolist():
        data = zin.read(item.filename)
        if item.filename in REPLACE:
            with open(REPLACE[item.filename], 'rb') as f:
                data = f.read()
        zout.writestr(item, data)
for _ in range(5):
    try:
        os.replace(TMP, DST)
        break
    except PermissionError:
        time.sleep(0.6)
else:
    raise SystemExit('✗ 仍无法写入 v24，请关闭 WPS 后重试')
print('✓ media 字节已替换:')
for k, v in REPLACE.items():
    print(f'   {k} ← {v}')

# ---------- 3. 验证 ----------
doc2 = Document(DST)
print('\n=== 验证 ===')
for p in doc2.paragraphs:
    for d in p._p.iter(qn('w:drawing')):
        blip = d.find('.//' + NS_A + 'blip')
        if blip is None:
            continue
        rId = blip.get(NS_R + 'embed')
        part = doc2.part.related_parts.get(rId)
        media = part.partname.split('/')[-1]
        if media in PLAN:
            ext = d.find('.//' + NS_WP + 'extent')
            cx = int(ext.get('cx')); cy = int(ext.get('cy'))
            blen = len(part.blob)
            print(f'   {media}: 显示 {cx/EMU:.2f}×{cy/EMU:.2f}cm | media字节 {blen/1024/1024:.2f}MB')
